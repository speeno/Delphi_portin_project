"""DEC-311 — 거래처거래원장 엑셀 = 화면(값·모양) + 일자별 출고 상세 반품/판매 컬럼 추가 (2026-09-24).

사용자: 「일자별 출고 도서 상세 내역에서 엑셀을 다운로드 받은 결과도 화면과 동일하게」,
「우측자료 "일자별 출고 상세보기" — 반품수량, 판매수량, 판매금액 확인불가 "추가요청"」.
"""

from __future__ import annotations

import io
from pathlib import Path
from unittest import TestCase, main

from openpyxl import load_workbook

from app.routers.export_table import TableExportColumn, _ALLOWED_NUM_FMTS
from app.services import masters_excel
from app.services.customer_txn_ledger_service import _detail_split

ROOT = Path(__file__).resolve().parents[1]
PAGE = ROOT / "도서물류관리프로그램" / "frontend" / "src" / "app" / "(app)" / "ledger" / "customer" / "page.tsx"


class DetailSplit(TestCase):
    def test_out_line(self) -> None:
        self.assertEqual(_detail_split(20, 374000, False),
                         {"out_qty": 20, "rtn_qty": 0, "sale_qty": 20, "sale_amt": 374000})

    def test_return_line_is_negative_and_counts_in_sale(self) -> None:
        # 반품은 음수 저장 관례 — 판매 = 출고 + 반품 이므로 그대로 판매에 가산(상단 전표 표와 같은 정의).
        self.assertEqual(_detail_split(-12, -570240, True),
                         {"out_qty": 0, "rtn_qty": -12, "sale_qty": -12, "sale_amt": -570240})


class WorkbookLooksLikeScreen(TestCase):
    def test_number_format_applies_to_numbers_only_and_blanks_stay_blank(self) -> None:
        cols = [("도서명", "gname"), ("공급률(%)", "grat1"), ("출고금액", "out_amt"), ("반품수량", "rtn_qty")]
        rows = [
            {"gname": "A", "grat1": 85, "out_amt": 3793020, "rtn_qty": ""},
            {"gname": "B", "grat1": 88, "out_amt": "", "rtn_qty": -12},
        ]
        data = masters_excel.build_list_workbook(
            sheet_title="t", columns=cols, rows=rows, number_formats=[None, None, "#,##0", "#,##0"],
        )
        ws = load_workbook(io.BytesIO(data)).active
        self.assertEqual(ws["C2"].value, 3793020, "값은 숫자 그대로(엑셀 합계 가능)")
        self.assertEqual(ws["C2"].number_format, "#,##0", "화면처럼 천 단위 구분")
        self.assertEqual(ws["D3"].number_format, "#,##0")
        self.assertIn(ws["D2"].value, (None, ""), "화면 빈칸 = 엑셀 빈칸")
        self.assertEqual(ws["B2"].number_format, "General", "공급률은 기본 서식(소수 반올림 방지)")

    def test_router_accepts_only_whitelisted_formats(self) -> None:
        col = TableExportColumn.model_validate({"key": "a", "label": "A", "numFmt": "#,##0"})
        self.assertEqual(col.num_fmt, "#,##0")
        self.assertIn("#,##0", _ALLOWED_NUM_FMTS)
        self.assertNotIn('[Red]"x"@', _ALLOWED_NUM_FMTS)


class PageStatic(TestCase):
    def setUp(self) -> None:
        self.src = PAGE.read_text(encoding="utf-8").replace("\r\n", "\n")

    def test_detail_grid_has_return_and_sale_columns(self) -> None:
        block = self.src.split("const DETAIL_COLUMNS")[1].split("];")[0]
        labels = ["출고수량", "출고금액", "반품수량", "반품금액", "판매수량", "판매금액", "미수"]
        pos = [block.index(f'label: "{l}"') for l in labels]
        self.assertEqual(pos, sorted(pos), "상단 전표 표와 같은 순서")
        self.assertNotIn('label: "수량"', block, "부호 있는 순수량(=판매수량) 중복 컬럼 제거")

    def test_export_blanks_zero_and_sends_comma_format(self) -> None:
        self.assertIn("function asScreenRow(", self.src)
        self.assertIn('numFmt: "#,##0"', self.src)
        for k in ("out_qty", "rtn_qty", "sale_qty", "sale_amt"):
            self.assertIn(f'"{k}"', self.src.split("const BLANK_ZERO_KEYS")[1].split("]);")[0])
        self.assertIn("withSplit", self.src, "구 백엔드 응답도 분해 보정")


if __name__ == "__main__":
    main()
