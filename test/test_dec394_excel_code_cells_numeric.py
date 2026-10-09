"""DEC-394 — 엑셀 다운로드의 코드 칸이 VLOOKUP 등 수식에 걸리도록 숫자 셀 + 자릿수 0 서식.

보고(2026-10-09, [공통]): 다운로드한 「거래처 코드」 가 VLOOKUP 에 반응하지 않음 — 텍스트 셀
「1234」「00039」 는 사용자가 친 1234 · 00039(엑셀이 숫자로 바꿈)와 일치하지 않는다.
"""

from __future__ import annotations

from io import BytesIO
from unittest import TestCase

from openpyxl import load_workbook

from app.services import masters_excel


def _sheet(columns, rows):
    data = masters_excel.build_list_workbook(sheet_title="목록", columns=columns, rows=rows)
    return load_workbook(BytesIO(data)).active


class CodeCellsNumericTests(TestCase):
    def test_digit_codes_become_numbers_with_zero_padded_display(self) -> None:
        ws = _sheet(
            [("거래처코드", "gcode"), ("거래처명", "gname")],
            [{"gcode": "00039", "gname": "가"}, {"gcode": "1234", "gname": "나"}, {"gcode": "A0001", "gname": "다"}],
        )
        self.assertEqual([ws.cell(r, 1).value for r in (2, 3, 4)], [39, 1234, "A0001"])
        self.assertEqual(ws.cell(2, 1).number_format, "00000")  # 화면 그대로 「00039」
        self.assertEqual(ws.cell(3, 1).number_format, "0000")
        self.assertEqual(ws.cell(4, 1).number_format, "General")  # 영문 섞인 코드는 텍스트 그대로

    def test_code_headers_only(self) -> None:
        """「코드」「도서코드2」 는 대상, 우편번호 · 이름 등 다른 칸의 숫자 문자열은 그대로."""
        ws = _sheet(
            [("코드", "gcode"), ("도서코드2", "ocode"), ("우편번호", "gpost"), ("거래처명", "gname")],
            [{"gcode": "0012", "ocode": "077", "gpost": "01000", "gname": "0001"}],
        )
        self.assertEqual([c.value for c in ws[2]], [12, 77, "01000", "0001"])

    def test_colliding_codes_keep_whole_column_text(self) -> None:
        """같은 테넌트에 「039」·「0039」 가 따로 있으면(실측 38개 테넌트) 숫자로는 구별 불가 → 칸 전체 텍스트."""
        ws = _sheet(
            [("코드", "gcode"), ("도서코드", "bcode")],
            [{"gcode": "039", "bcode": "100"}, {"gcode": "0039", "bcode": "00200"}, {"gcode": "1234", "bcode": "300"}],
        )
        self.assertEqual([ws.cell(r, 1).value for r in (2, 3, 4)], ["039", "0039", "1234"])
        self.assertEqual([ws.cell(r, 2).value for r in (2, 3, 4)], [100, 200, 300])  # 다른 칸은 영향 없음

    def test_over_15_digits_stays_text(self) -> None:
        ws = _sheet([("코드", "gcode")], [{"gcode": "1234567890123456"}])
        self.assertEqual(ws.cell(2, 1).value, "1234567890123456")

    def test_generic_table_export_route_applies_too(self) -> None:
        from fastapi.testclient import TestClient

        from app.core.deps import get_user_context
        from app.main import app

        app.dependency_overrides[get_user_context] = lambda: {"user_id": "u", "server_id": "remote_1", "hcode": "5019"}
        try:
            res = TestClient(app).post(
                "/api/v1/export/table-xlsx",
                json={"columns": [{"key": "customer_code", "label": "거래처코드"}],
                      "rows": [{"customer_code": "00039"}]},
            )
        finally:
            app.dependency_overrides.pop(get_user_context, None)
        self.assertEqual(res.status_code, 200, res.text[:200])
        cell = load_workbook(BytesIO(res.content)).active.cell(2, 1)
        self.assertEqual((cell.value, cell.number_format), (39, "00000"))


class ImportRestoresCodeTests(TestCase):
    """다운로드 → 수정 → 업로드(역반영) 왕복에서 「00039」 가 「39」 로 바뀌면 다른 거래처를 고친다."""

    def test_round_trip_restores_leading_zeros(self) -> None:
        cols = masters_excel.select_customer_columns(["gcode", "gname"])
        data = masters_excel.build_list_workbook(
            sheet_title="거래처", columns=cols,
            rows=[{"gcode": "00039", "gname": "가"}, {"gcode": "1234", "gname": "나"}, {"gcode": "A0001", "gname": "다"}],
        )
        parsed = masters_excel.parse_master_xlsx(
            data, pk=masters_excel.CUSTOMER_IMPORT_PK, field_map=masters_excel.CUSTOMER_IMPORT_MAP,
            pk_aliases=masters_excel.CUSTOMER_IMPORT_PK_ALIASES,
        )
        self.assertEqual([r["gcode"] for r in parsed["rows"]], ["00039", "1234", "A0001"])

    def test_code_text_helper(self) -> None:
        self.assertEqual(masters_excel._code_text(39, "00000"), "00039")
        self.assertEqual(masters_excel._code_text(39.0, "00000"), "00039")
        self.assertEqual(masters_excel._code_text(39, "General"), "39")  # 사용자가 직접 친 숫자
        self.assertEqual(masters_excel._code_text(" A1 ", None), "A1")
